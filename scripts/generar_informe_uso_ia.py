"""Genera el informe de uso de inteligencia artificial del trabajo del PETI.

Uso:
    python scripts/generar_informe_uso_ia.py [salida.docx]

Reproduce el formato de la declaración de uso de IA de la Universidad Santiago de Cali
(Arial 10, interlineado 21 pt, tabla de aspectos con filas alternas sombreadas) e inserta
la firma del líder de equipo desde docs/privado/firma-lider-equipo.png. Requiere python-docx.

Antes de generarlo, reemplazar todos los [COMPLETAR] de las constantes de texto con los datos
reales del PETI (cifras de brechas, proyectos, indicadores y riesgos iguales a las del documento).
Si la universidad o la circular son otras, ajustar ENTRADA.
"""
import os
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIRMA = os.path.join(RAIZ, "docs", "privado", "firma-lider-equipo.png")
SALIDA = os.path.join(RAIZ, "docs", "Informe_uso_IA_PETI.docx")

GRIS = "D9D9D9"
TINTA = RGBColor(0x0F, 0x11, 0x15)
FECHA = "[COMPLETAR: día de mes de año]"

TITULO = "Informe de uso de inteligencia artificial"
SUBTITULO = [
    "Plan Estratégico de Tecnologías de la Información (PETI) [COMPLETAR: periodo]",
    "[COMPLETAR: organización — alcance]",
    "Trabajo académico de Gobierno de TI",
]

RESUMEN_INTRO = (
    "Este informe acompaña al PETI que se entrega y declara cómo se usó la inteligencia artificial "
    "para construirlo. El trabajo se resume así."
)

RESUMEN = [
    ("Nombre del proyecto",
     "Plan Estratégico de Tecnologías de la Información (PETI) [COMPLETAR: periodo] de "
     "[COMPLETAR: organización]. Trabajo académico de Gobierno de TI; el documento corresponde a la "
     "versión [COMPLETAR]."),
    ("Objetivo o propósito",
     "[COMPLETAR: el objetivo del §2.1 del PETI, con periodo y las brechas principales que cierra.]"),
    ("Alcance y fondo",
     "[COMPLETAR: qué es la organización, qué hecho reciente motiva el plan, qué unidades y dominios "
     "cubre y qué queda fuera (§2.2 del PETI).]"),
    ("Metodología o enfoque",
     "Guía MGGTI.GE.ES.03 v3.0 del MinTIC, con su plantilla de documento y sus herramientas de "
     "construcción [COMPLETAR si es privada: ', adaptadas a una empresa privada según lo que la propia "
     "guía permite']. Cada hecho de la situación actual lleva identificador de evidencia y una "
     "clasificación de pública, interna o restringida; lo que no la tiene se redacta como hipótesis y "
     "queda marcado."),
    ("Principales hallazgos o resultados",
     "El diagnóstico identificó [n] brechas entre la situación actual y la objetivo. El portafolio "
     "responde con [n] proyectos organizados en [n] fases, [n] indicadores de gestión y [n] riesgos "
     "con su respectivo tratamiento. [COMPLETAR: las dos o tres brechas que más pesan.]"),
    ("Impacto y beneficios",
     "[COMPLETAR: los efectos medibles del plan, iguales a los beneficios del §1 del PETI.]"),
    ("Conclusión y próximos pasos",
     "[COMPLETAR: versión entregada y próximos pasos: validar hipótesis con fuentes primarias, "
     "cerrar la auditoría de conformidad, aprobación por la instancia de gobierno.]"),
]

ENTRADA = [
    ("En cumplimiento de lo establecido en la ", {}),
    ("Circular Normativa DGI No. 002", {"bold": True}),
    (" de la Universidad Santiago de Cali sobre el uso responsable de la IA en los trabajos "
     "académicos, se declara lo siguiente:", {}),
]

TABLA_IA = [
    ("Herramientas de IA utilizadas",
     "Claude (Anthropic), mediante Claude Code, con una configuración propia guardada en el "
     "repositorio del trabajo."),
    ("Propósito del uso",
     "Lectura y extracción de la Guía MGGTI.GE.ES.03 v3.0 del MinTIC, de la plantilla PETI y de "
     "las herramientas de construcción; organización de los insumos de investigación sobre "
     "la organización; armado de los catálogos de brechas, iniciativas e indicadores y de la "
     "matriz de trazabilidad; verificación de las afirmaciones "
     "públicas contra su fuente primaria; auditoría de conformidad del documento contra la norma; "
     "sugerencias de redacción y generación del formato en Word."),
    ("Grado de intervención humana",
     "El equipo eligió la organización, el horizonte y el alcance del plan, y tomó las decisiones "
     "metodológicas: [COMPLETAR: p. ej. cómo se adapta a una empresa privada una guía escrita para "
     "entidades públicas, qué marco legal aplica, qué información de las entrevistas podía usarse y "
     "cómo], y con qué criterio se prioriza el portafolio. Cada numeral fue revisado y aprobado por "
     "el equipo; las respuestas de la herramienta se trataron como sugerencias."),
    ("Originalidad del contenido",
     "El PETI es un ejercicio propio del equipo sobre información pública de [COMPLETAR: "
     "organización] y sobre evidencia registrada con identificador. El texto se "
     "elaboró con asistencia de IA mediante sugerencias. Las afirmaciones que aún no tienen fuente "
     "verificada están marcadas en el documento."),
]

CONFIG_INTRO = (
    "La herramienta no se usó en su configuración estándar. El equipo armó en el repositorio una "
    "carpeta de configuración (.claude) que no contiene texto del PETI: fija cómo debe trabajar la "
    "herramienta dentro de este proyecto. Consta de cuatro partes."
)

CONFIG_TABLA = [
    ("Contexto del proyecto",
     "Un archivo de instrucciones permanentes que distingue qué documento es norma (la guía del "
     "MinTIC y sus anexos), cuál es referencia y cuál es insumo de investigación, y fija las "
     "convenciones de identificadores del plan."),
    ("Habilidades por tarea",
     "Instrucciones y scripts para cada tipo de trabajo: qué exige la norma en cada capítulo, "
     "criterios de estilo, registro y clasificación de fuentes, citación de marcos como COBIT o "
     "ISO/IEC 27001, y verificación de la cadena objetivo, brecha, proyecto, indicador y riesgo."),
    ("Revisiones independientes",
     "Tres revisores automáticos que leen el documento terminado y reportan hallazgos: uno "
     "contrasta el PETI con la guía y la plantilla, otro comprueba que cada hecho tenga evidencia "
     "registrada, y el tercero verifica en la web las afirmaciones públicas sobre la empresa."),
    ("Reglas y permisos",
     "Restricciones sobre qué archivos pueden modificarse y cuáles no se versionan. Los insumos de "
     "entrevistas y los documentos internos quedan excluidos del control de "
     "versiones y no se copian a ningún archivo compartido."),
]

CONFIG_CIERRE = (
    "El efecto de esa configuración es que la metodología del MinTIC queda a la vista en cada "
    "paso y que las afirmaciones del plan se pueden rastrear hasta la fuente que las respalda."
)

DECLARACION = [
    ("Declaro que", {"bold": True}),
    (" asumo la responsabilidad total sobre la originalidad y veracidad de la información "
     "presentada en este trabajo. Se utilizó una herramienta de inteligencia artificial para "
     "extraer la metodología del MinTIC y sus formatos, organizar los insumos de investigación, "
     "armar los catálogos y las matrices del plan, verificar las fuentes públicas, auditar la "
     "conformidad del documento y proponer sugerencias de redacción y de formato. El equipo "
     "definió el alcance del trabajo, tomó las decisiones metodológicas y revisó y aprobó el "
     "contenido final.", {}),
]


def fuente(run, tam=10, negrita=False, cursiva=False):
    run.font.name = "Arial"
    run.font.size = Pt(tam)
    run.font.color.rgb = TINTA
    run.bold = negrita
    run.italic = cursiva
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atributo), "Arial")
    return run


def parrafo(doc, partes, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, tam=10, antes=12, despues=12,
            interlineado=21):
    p = doc.add_paragraph()
    p.alignment = alineacion
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    if interlineado:
        pf.line_spacing = Pt(interlineado)
    if isinstance(partes, str):
        partes = [(partes, {})]
    for texto, formato in partes:
        fuente(p.add_run(texto), tam, formato.get("bold", False), formato.get("italic", False))
    return p


def titulo_seccion(doc, texto, antes=18, despues=6):
    p = parrafo(doc, texto, WD_ALIGN_PARAGRAPH.LEFT, antes=antes, despues=despues)
    for run in p.runs:
        run.bold = True
    return p


def sombrear(celda, color):
    tc = celda._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def bordes_horizontales(tabla):
    """Solo líneas superior, inferior y bajo el encabezado; sin líneas verticales."""
    tbl_pr = tabla._element.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado, ancho in (("top", 8), ("bottom", 8), ("insideH", 0), ("insideV", 0),
                        ("left", 0), ("right", 0)):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single" if ancho else "none")
        b.set(qn("w:sz"), str(ancho))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "0F1115" if ancho else "auto")
        bordes.append(b)
    tbl_pr.append(bordes)


def borde_celda(celda, lado, ancho=6):
    tc_pr = celda._element.get_or_add_tcPr()
    bordes = tc_pr.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        tc_pr.append(bordes)
    b = OxmlElement(f"w:{lado}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(ancho))
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), "0F1115")
    bordes.append(b)


def no_partir(fila):
    """Evita que una fila se corte entre dos páginas."""
    fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def celda_texto(celda, texto, negrita=False, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = celda.paragraphs[0]
    p.alignment = alineacion
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    pf.line_spacing = Pt(14)
    fuente(p.add_run(texto), 10, negrita)


def tabla_aspectos(doc, filas, encabezado=("Aspecto", "Declaración")):
    tabla = doc.add_table(rows=1, cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    bordes_horizontales(tabla)
    cabecera = tabla.rows[0].cells
    for i, texto in enumerate(encabezado):
        celda_texto(cabecera[i], texto, negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
        borde_celda(cabecera[i], "bottom", 8)
    tabla.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, (aspecto, declaracion) in enumerate(filas):
        celdas = tabla.add_row().cells
        celda_texto(celdas[0], aspecto, negrita=True, alineacion=WD_ALIGN_PARAGRAPH.LEFT)
        celda_texto(celdas[1], declaracion)
        if i % 2 == 0:
            for c in celdas:
                sombrear(c, GRIS)
    for fila in tabla.rows:
        fila.cells[0].width = Cm(5.2)
        fila.cells[1].width = Cm(10.4)
        no_partir(fila)
    return tabla


def bloque_firma(doc, etiqueta, ruta_firma):
    """Etiqueta a la izquierda y la firma sobre una línea, como en el formato original."""
    tabla = doc.add_table(rows=1, cols=2)
    tabla.autofit = False
    celdas = tabla.rows[0].cells
    celdas[0].width = Cm(5.6)
    celdas[1].width = Cm(6.5)
    celda_texto(celdas[0], etiqueta, negrita=True, alineacion=WD_ALIGN_PARAGRAPH.LEFT)
    p = celdas[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if os.path.exists(ruta_firma):
        p.add_run().add_picture(ruta_firma, width=Cm(4.4))
    else:
        fuente(p.add_run(" "), 10)
        print(f"AVISO: no se encontró la firma en {ruta_firma}; la línea queda en blanco.")
    borde_celda(celdas[1], "bottom", 6)
    return tabla


def main(salida):
    doc = docx.Document()
    for seccion in doc.sections:
        seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)  # carta
        seccion.left_margin = seccion.right_margin = Cm(3.0)
        seccion.top_margin = seccion.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    parrafo(doc, TITULO, WD_ALIGN_PARAGRAPH.CENTER, tam=13, antes=0, despues=6)
    for i, linea in enumerate(SUBTITULO):
        parrafo(doc, linea, WD_ALIGN_PARAGRAPH.CENTER, tam=10, antes=0,
                despues=12 if i == len(SUBTITULO) - 1 else 0, interlineado=14)
    for run in doc.paragraphs[0].runs:
        run.bold = True

    titulo_seccion(doc, "Resumen ejecutivo", antes=6)
    parrafo(doc, RESUMEN_INTRO)
    tabla_aspectos(doc, RESUMEN, encabezado=("Componente", "Contenido"))

    titulo_seccion(doc, "Declaración de uso de inteligencia artificial")
    parrafo(doc, ENTRADA)
    tabla_aspectos(doc, TABLA_IA)

    titulo_seccion(doc, "Configuración de la herramienta")
    parrafo(doc, CONFIG_INTRO)
    tabla_aspectos(doc, CONFIG_TABLA, encabezado=("Parte", "Qué define"))
    parrafo(doc, CONFIG_CIERRE)

    parrafo(doc, DECLARACION, antes=18)
    bloque_firma(doc, "Firma del líder de equipo:", FIRMA)
    parrafo(doc, [("Fecha: ", {"bold": True}), (FECHA, {})], WD_ALIGN_PARAGRAPH.LEFT, antes=14)

    doc.save(salida)
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else SALIDA)
