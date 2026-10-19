"""Genera el Word del PETI editando una copia de la Plantilla PETI del MinTIC (Anexo 2).

Uso:
    python .claude/skills/guia-peti/scripts/plantilla_peti.py <entrada.md> <salida.docx>

A diferencia de md_a_docx.py, no parte de un documento vacío: conserva la portada, las páginas
separadoras de capítulo, los títulos, los estilos y el formato de las tablas de la plantilla, y
reemplaza los textos guía y los ejemplos por el contenido del Markdown.

Reglas:
  - Cada título del Markdown ('##' a '#####') se busca en la plantilla, en orden, por su texto.
    Si coincide, el contenido va debajo de ese título; si no, se crea un título nuevo con el
    estilo de la plantilla (numerado "N.N" -> Título 3; sin número -> Título 2 si es '##', Título 4 si es más profundo).
  - La numeración de la plantilla no se cambia; RENOMBRAR solo corrige erratas de texto.
  - Una línea "Tabla: descripción {Tnn}" antes de una tabla indica qué tabla de la plantilla se
    clona (índice en el orden del documento original) y el texto de su título ("Tabla N ...").
  - Se eliminan los párrafos guía, los ejemplos, las tablas de ejemplo, los íconos de los
    recuadros de guía y las ilustraciones de ejemplo. Se conservan los saltos de página y de
    sección y las imágenes decorativas de las páginas separadoras.
  - Lo que va antes del primer '##' del Markdown (título y versión) se escribe en la portada.
Requiere python-docx.
"""
import copy
import os
import re
import sys
import unicodedata

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
PLANTILLA = os.path.join(RAIZ, "docs", "Anexo 2. Plantilla PETI.docx")

RENOMBRAR = {
    # Solo se corrige la errata del título; la numeración de la plantilla se conserva tal cual
    "7. Situación deseada u objetivo": "7. Situación de seada u Objetivo",
}
# Tablas con dos filas de encabezado en la plantilla (grupo de columnas + nombres de columna)
DOS_ENCABEZADOS = {66, 100}
# Fichas cuya tercera columna es una sola celda combinada de acciones (servicios y sistemas objetivo)
FICHAS_CON_ACCIONES = {69, 104}
# Tablas cuya primera columna (categoría) se combina verticalmente (capacidades de TI)
AGRUPAR_PRIMERA_COLUMNA = {42, 68}
# Desde este número de columnas la tabla va en una página horizontal
HORIZONTAL = 10
# Tamaño (cm, redondeado) de las imágenes decorativas de las páginas separadoras
DECORATIVAS = {(7, 28), (6, 6), (17, 16), (14, 11)}
ESTILOS_TITULO = {"Heading 1", "Heading 2", "Heading 3", "Heading 4", "Subtitle", "No Spacing"}


def norm(texto):
    texto = unicodedata.normalize("NFD", texto.lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    # Sin espacios: en la plantilla, el número y el título a veces van separados por un w:tab
    return re.sub(r"\s+", "", texto).rstrip(":")


def texto_de(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


# ---------- lectura del Markdown ----------

def leer_markdown(ruta):
    lineas = open(ruta, encoding="utf8").read().splitlines()
    portada, bloques, tabla, pendiente = [], [], [], None
    en_cuerpo = False
    for linea in lineas + [""]:
        s = linea.rstrip()
        if not en_cuerpo:
            if s.startswith("## "):
                en_cuerpo = True
            else:
                if s.strip():
                    portada.append(s)
                continue
        if s.strip().startswith("|"):
            tabla.append([c.strip() for c in s.strip().strip("|").split("|")])
            continue
        if tabla:
            filas = [f for f in tabla if not all(re.fullmatch(r":?-+:?", c or "-") for c in f)]
            desc, indice = pendiente or (None, 101)
            bloques.append(("tabla", indice, desc, filas))
            tabla, pendiente = [], None
        if not s.strip():
            continue
        m = re.match(r"^(#{2,5})\s+(.*)", s)
        mt = re.match(r"^Tabla:\s*(.*?)\s*\{T(\d+)\}\s*$", s)
        mf = re.match(r"^Figura:\s*(.*?)\s*\{img:(.+?)\}\s*$", s)
        if mf:
            bloques.append(("imagen", mf.group(1), mf.group(2)))
        elif m:
            bloques.append(("titulo", len(m.group(1)), m.group(2).strip()))
        elif mt:
            pendiente = (mt.group(1), int(mt.group(2)))
        elif re.match(r"^\s*[-*]\s+", s):
            bloques.append(("vineta", re.sub(r"^\s*[-*]\s+", "", s)))
        else:
            bloques.append(("parrafo", s[2:] if s.startswith("> ") else s))
    return portada, bloques


# ---------- utilidades de párrafos y tablas ----------

def rpr_de(el):
    for r in el.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            return copy.deepcopy(rpr)
    return None


def agregar_runs(p, texto, rpr=None, tam=None):
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not trozo:
            continue
        r = OxmlElement("w:r")
        props = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
        for etiqueta in ("w:b", "w:i"):
            viejo = props.find(qn(etiqueta))
            if viejo is not None and trozo.startswith("*"):
                props.remove(viejo)
        if trozo.startswith("**"):
            props.append(OxmlElement("w:b")); trozo = trozo[2:-2]
        elif trozo.startswith("*"):
            props.append(OxmlElement("w:i")); trozo = trozo[1:-1]
        if tam:
            for etiqueta in ("w:sz", "w:szCs"):
                viejo = props.find(qn(etiqueta))
                if viejo is not None:
                    props.remove(viejo)
                e = OxmlElement(etiqueta); e.set(qn("w:val"), str(tam * 2)); props.append(e)
        if len(props):
            r.append(props)
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = trozo
        r.append(t)
        p.append(r)


def vaciar(p):
    for hijo in list(p):
        if hijo.tag != qn("w:pPr"):
            p.remove(hijo)


def poner_texto(p, texto, tam=None):
    rpr = rpr_de(p)
    vaciar(p)
    agregar_runs(p, texto, rpr, tam)


def parrafo_nuevo(estilo_id=None, ppr=None):
    p = OxmlElement("w:p")
    if ppr is not None:
        p.append(copy.deepcopy(ppr))
    elif estilo_id:
        pp = OxmlElement("w:pPr"); ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), estilo_id)
        pp.append(ps); p.append(pp)
    return p


def llenar_celda(tc, texto, tam=None):
    ps = tc.findall(qn("w:p"))
    for extra in ps[1:]:
        tc.remove(extra)
    for tbl in tc.findall(qn("w:tbl")):
        tc.remove(tbl)
    poner_texto(ps[0], texto, tam)


def ajustar_fila(tr, ncol, ancho, quitar_combinadas=True):
    tcs = tr.findall(qn("w:tc"))
    if quitar_combinadas:
        for tc, w in zip(tcs, ancho):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is not None:
                for etiqueta in ("w:gridSpan", "w:vMerge", "w:hMerge"):
                    for e in tcpr.findall(qn(etiqueta)):
                        tcpr.remove(e)
    while len(tcs) > ncol:
        tr.remove(tcs.pop())
    while len(tcs) < ncol:
        nuevo = copy.deepcopy(tcs[-1]); tcs[-1].addnext(nuevo); tcs.append(nuevo)
    if quitar_combinadas:
        for tc, w in zip(tcs, ancho):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr"); tc.insert(0, tcpr)
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW"); tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(w)); tcw.set(qn("w:type"), "dxa")
    return tcs


def clonar_tabla(modelo, indice, filas, ancho_total=None):
    tbl = copy.deepcopy(modelo)
    ncol = max(len(f) for f in filas)
    grid = tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    total = sum(int(c.get(qn("w:w")) or 0) for c in cols) or 9000
    mismo_grid = len(cols) == ncol  # mismas columnas: se conservan las celdas combinadas del encabezado
    total = ancho_total or total
    trs = tbl.findall(qn("w:tr"))
    enc = 2 if indice in DOS_ENCABEZADOS and mismo_grid and len(trs) > 2 else 1
    prototipo = copy.deepcopy(trs[enc] if len(trs) > enc else trs[-1])
    if indice in DOS_ENCABEZADOS:
        # En estas tablas la fila de ejemplo trae el color del encabezado: los datos van sin relleno
        for etiqueta in ("w:shd", "w:color"):
            for e in list(prototipo.iter(qn(etiqueta))):
                e.getparent().remove(e)
    for tr in trs[enc:]:
        tbl.remove(tr)
    # Ancho de cada columna según el texto que lleva (con un mínimo), no igual para todas
    pesos = []
    for j in range(ncol):
        largos = [len(f[j]) for f in filas if j < len(f)]
        pesos.append(min(max(sum(largos) / len(largos), 6), 90))
    anchos = [max(int(total * w / sum(pesos)), 600) for w in pesos]
    if mismo_grid and not ancho_total:
        anchos = [int(c.get(qn("w:w")) or total // ncol) for c in cols]
    if ancho_total:
        tblpr = tbl.find(qn("w:tblPr"))
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW"); tblpr.append(tblw)
        tblw.set(qn("w:w"), str(total)); tblw.set(qn("w:type"), "dxa")
    if not mismo_grid or ancho_total:
        for c in cols:
            grid.remove(c)
        for w in anchos:
            g = OxmlElement("w:gridCol"); g.set(qn("w:w"), str(w)); grid.append(g)
    ancho = anchos
    tam = 8 if ncol >= 7 else None
    # encabezado: con dos filas, la primera (grupos de columnas) queda como en la plantilla
    fila_enc = tbl.findall(qn("w:tr"))[enc - 1]
    for tc, texto in zip(ajustar_fila(fila_enc, ncol, ancho, quitar_combinadas=not mismo_grid), filas[0]):
        llenar_celda(tc, texto, tam)
    for fila in filas[1:]:
        tr = copy.deepcopy(prototipo)
        tcs = ajustar_fila(tr, ncol, ancho, quitar_combinadas=True)
        for j, tc in enumerate(tcs):
            llenar_celda(tc, fila[j] if j < len(fila) else "", tam)
        tbl.append(tr)
    trs = tbl.findall(qn("w:tr"))
    if indice in FICHAS_CON_ACCIONES and ncol == 3 and enc == 1 and len(trs) > 2:
        # Como en la plantilla: la tercera columna es una sola celda combinada con el título y las acciones
        llenar_celda(trs[0].findall(qn("w:tc"))[2], "", tam)
        acciones = [f"{f[0]}: {f[2]}" for f in filas[2:] if len(f) > 2 and f[2] and f[2].lower() != "se mantiene"]
        titulo = filas[1][2] if len(filas[1]) > 2 and filas[1][2] and filas[1][2].lower() != "se mantiene" else ""
        if titulo:
            acciones.insert(0, f"{filas[1][0]}: {titulo}")
        combinar_vertical(trs, 2, 1, len(trs) - 1,
                          [f"**{filas[0][2]}**"] + (acciones or ["Se mantiene en todos sus atributos."]), tam)
    if indice in AGRUPAR_PRIMERA_COLUMNA and enc == 1:
        # Como en la plantilla: la categoría se combina verticalmente sobre sus filas
        r = 1
        while r < len(trs):
            s = r
            while s + 1 < len(trs) and filas[s + 1][0] == filas[r][0]:
                s += 1
            if s > r:
                combinar_vertical(trs, 0, r, s, [filas[r][0]], tam)
            r = s + 1
    return tbl


def combinar_vertical(trs, col, ini, fin, lineas, tam=None):
    """Combina verticalmente la columna col de las filas ini..fin y escribe las líneas en la primera celda."""
    for r in range(ini, fin + 1):
        tc = trs[r].findall(qn("w:tc"))[col]
        tcpr = tc.find(qn("w:tcPr"))
        if tcpr is None:
            tcpr = OxmlElement("w:tcPr"); tc.insert(0, tcpr)
        for e in tcpr.findall(qn("w:vMerge")):
            tcpr.remove(e)
        vm = OxmlElement("w:vMerge")
        if r == ini:
            vm.set(qn("w:val"), "restart")
        tcpr.append(vm)
        if r == ini:
            llenar_celda(tc, lineas[0], tam)
            base = tc.findall(qn("w:p"))[0]
            for linea in lineas[1:]:
                p = copy.deepcopy(base)
                poner_texto(p, linea, tam)
                # el título va como en la plantilla (negrita, centrado); cada acción, normal y a la izquierda
                for e in list(p.iter(qn("w:b"))) + list(p.iter(qn("w:bCs"))):
                    e.getparent().remove(e)
                ppr = p.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
                for e in ppr.findall(qn("w:jc")):
                    ppr.remove(e)
                jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "left"); ppr.append(jc)
                tc.append(p)
        else:
            llenar_celda(tc, "", tam)


# ---------- conversión ----------

def estilo(el, doc):
    if el.tag != qn("w:p"):
        return None
    try:
        return Paragraph(el, doc._body).style.name
    except Exception:
        return None


def tamano(dr):
    ext = dr[0].find(qn("wp:extent"))
    return (round(int(ext.get("cx")) / 360000), round(int(ext.get("cy")) / 360000))


def es_decorativa(dr):
    return dr[0].tag == qn("wp:anchor") and tamano(dr) in DECORATIVAS and not texto_de(dr)


def estructural(el):
    """Párrafo que se conserva sin texto: salto de sección o de página, o imagen decorativa."""
    if el.tag != qn("w:p"):
        return False
    ppr = el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
        return True
    if any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br"))):
        return True
    return any(es_decorativa(dr) for dr in el.iter(qn("w:drawing")))


def limpiar_estructural(el):
    for r in list(el.iter(qn("w:r"))):
        util = any(br.get(qn("w:type")) == "page" for br in r.iter(qn("w:br"))) or \
            any(es_decorativa(dr) for dr in r.iter(qn("w:drawing")))
        if not util and r.getparent() is not None:
            r.getparent().remove(r)
    for h in list(el.findall(qn("w:hyperlink"))):
        el.remove(h)


def portada(doc, lineas_md):
    sdt = doc.element.body[0]
    ps = list(sdt.iter(qn("w:p")))
    version = next((l for l in lineas_md if l.lower().startswith("versión")), "Versión")
    m = re.search(r"Versión\s+([\d.]+)", version)
    numero = m.group(1) if m else ""
    cambios = {
        3: "2027-2030", 5: "2027-2030",
        7: "Valor + S.A.S. — enero de 2027 a diciembre de 2030",
        8: "Valor + S.A.S. — enero de 2027 a diciembre de 2030",
        9: "Valor + S.A.S.",
        10: "Plan Estratégico de Tecnologías de la Información (PETI) 2027-2030",
        11: "Alcance: toda la organización: TI interna y catálogo de servicios de TI",
        12: "Metodología: Guía MGGTI.GE.ES.03 v3.0 del MinTIC",
        14: "Equipo de trabajo",
        15: "Hannyer Arce Ararat – Líder y coordinador",
        16: "Kevin Steven Tabares Blanco – Contexto y normatividad",
        17: "Raul Guillermo Orobio Ospina – Diagnóstico de estrategia, gobierno e información; evidencia",
        18: "Ever Jonatan Delgado Chaguendo – Diagnóstico de sistemas, infraestructura, uso y seguridad",
        20: "Juan Jose Muñoz – Situación objetivo y brechas",
        21: "Luis Fernando Narvaes Marin – Portafolio, indicadores y riesgos",
        23: "Equipo académico del PETI, en el rol del área de TI de Valor+",
        31: f"Versión {numero}",
        32: "Octubre de 2026",
        33: "Propuesta sujeta a la aprobación del Comité Institucional de Gestión y Desempeño",
        37: "Nota sobre este documento",
        39: "Este es un documento académico de Gobierno de TI: un equipo de estudiantes asume el rol del área de "
            "TI de Valor + S.A.S. y formula su Plan Estratégico de Tecnologías de la Información con la metodología "
            "del MinTIC. No es un documento oficial de Valor+.",
        40: "Clasificación: interno. El plan se construyó solo con fuentes públicas; cada dato cita el documento del "
            "que sale y la lista completa está en Referencias Bibliográficas. Lo que no tiene fuente está marcado "
            "como hipótesis o propuesta.",
        41: "La estructura sigue la Plantilla PETI del MinTIC (Anexo 2): sus textos guía y ejemplos se reemplazaron "
            "por el contenido del plan. Todo lo que se propone depende de la aprobación del Comité Institucional "
            "de Gestión y Desempeño.",
        43: "", 44: "",
        49: "Octubre 2026", 50: "Octubre 2026",
    }
    for i, texto in cambios.items():
        poner_texto(ps[i], texto)
    # Logos del Gobierno (portada) y cuadro de convenciones de la guía (página de bienvenida)
    for i in (0, 45):
        for dr in list(ps[i].iter(qn("w:drawing"))):
            blips = [b.get(qn("r:embed")) for b in dr.iter(qn("a:blip"))]
            nombres = [doc.part.related_parts[r].partname for r in blips]
            if i == 45 or "/word/media/image1.png" in nombres:
                r = dr.getparent()
                r.getparent().remove(r)
    # Se quita la página "Nota sobre este documento" (decisión del usuario, 6 de octubre de 2026)
    for p in ps[37:47]:
        padre = p.getparent()
        if padre.tag == qn("w:tc") and len(padre.findall(qn("w:p"))) == 1:
            continue
        padre.remove(p)
    return ps[47]  # "Introducción", título dentro de la portada


def leyenda_con_numero(modelo, numero, desc):
    """Copia el título de tabla o ilustración de la plantilla con su campo SEQ y cambia número y texto."""
    cap = copy.deepcopy(modelo)
    runs = [r for r in cap if r.tag == qn("w:r")]
    fin = max(i for i, r in enumerate(runs)
              if any(f.get(qn("w:fldCharType")) == "end" for f in r.iter(qn("w:fldChar"))))
    sep = max(i for i, r in enumerate(runs)
              if any(f.get(qn("w:fldCharType")) == "separate" for f in r.iter(qn("w:fldChar"))))
    for r in runs[sep + 1:fin]:
        for t in r.iter(qn("w:t")):
            t.text = str(numero)
    rpr = rpr_de(runs[0])
    for r in runs[fin + 1:]:
        cap.remove(r)
    agregar_runs(cap, " " + desc, rpr)
    return cap


def util_horizontal(sect):
    """Ancho útil (dxa) de la página girada: el alto de la hoja menos los márgenes laterales."""
    pg, mar = sect.find(qn("w:pgSz")), sect.find(qn("w:pgMar"))
    return int(pg.get(qn("w:h"))) - int(mar.get(qn("w:left"))) - int(mar.get(qn("w:right"))) - 200


def convertir(entrada, salida):
    lineas_portada, bloques = leer_markdown(entrada)
    doc = docx.Document(PLANTILLA)
    body = doc.element.body
    # Tablas de la plantilla en el orden del documento (las del cuerpo; no las de la portada)
    modelos = [copy.deepcopy(t) for t in body if t.tag == qn("w:tbl")]

    ids = {}
    for el in body.iter(qn("w:p")):
        nombre = estilo(el, doc)
        ppr = el.find(qn("w:pPr"))
        if nombre and ppr is not None and ppr.find(qn("w:pStyle")) is not None:
            ids.setdefault(nombre, ppr.find(qn("w:pStyle")).get(qn("w:val")))
    vineta_ppr = next(copy.deepcopy(el.find(qn("w:pPr"))) for el in body
                      if estilo(el, doc) == "List Paragraph" and el.find(qn("w:pPr")).find(qn("w:numPr")) is not None)
    leyenda = next(el for el in body if estilo(el, doc) == "Caption"
                   and any("SEQ Tabla" in (x.text or "") for x in el.iter(qn("w:instrText"))))
    leyenda = copy.deepcopy(leyenda)
    leyenda_fig = next((copy.deepcopy(el) for el in body if estilo(el, doc) == "Caption"
                        and any("SEQ Ilustraci" in (x.text or "") for x in el.iter(qn("w:instrText")))), None)

    intro = portada(doc, lineas_portada)
    hijos = [el for el in body if el.tag != qn("w:sectPr")]
    fin_cuerpo = body.find(qn("w:sectPr"))

    # Títulos candidatos de la plantilla, en orden
    candidatos = [(0, intro)] + [(i, el) for i, el in enumerate(hijos) if i > 0 and estilo(el, doc) in ESTILOS_TITULO
                                 and texto_de(el).strip()]
    titulos_md = [b for b in bloques if b[0] == "titulo"]
    asignados, puntero, nuevos = {}, 0, []
    for k, (_, nivel, texto) in enumerate(titulos_md):
        buscado = norm(RENOMBRAR.get(texto, texto))
        for j in range(puntero, len(candidatos)):
            if norm(texto_de(candidatos[j][1])) == buscado:
                asignados[k] = candidatos[j]; puntero = j + 1
                break
        else:
            nuevos.append(texto)
    usados = {id(el) for _, el in asignados.values()}
    posiciones = sorted(i for i, _ in asignados.values())

    # Qué se conserva
    borrar = []
    tras_decoracion = False  # los párrafos vacíos que siguen a una ilustración la separan del texto
    for i, el in enumerate(hijos):
        decorado = el.tag == qn("w:p") and any(es_decorativa(dr) for dr in el.iter(qn("w:drawing")))
        vacio = el.tag == qn("w:p") and not texto_de(el).strip() and not list(el.iter(qn("w:drawing")))
        if i == 0 or id(el) in usados:
            if i > 0:
                # Los íconos de los recuadros de guía también van anclados a algunos títulos
                for dr in list(el.iter(qn("w:drawing"))):
                    if not es_decorativa(dr):
                        r = dr.getparent()
                        while r is not None and r.tag != qn("w:r"):
                            r = r.getparent()
                        if r is not None and r.getparent() is not None:
                            r.getparent().remove(r)
            tras_decoracion = decorado
            continue
        if estructural(el):
            limpiar_estructural(el)
            tras_decoracion = decorado or tras_decoracion
        elif vacio and tras_decoracion:
            pass
        else:
            borrar.append(el)
            tras_decoracion = False
    # Un salto de página solo se conserva si lo que sigue abre capítulo o sección, y solo uno por tramo
    anclas_mayores = {id(el) for _, el in asignados.values()
                      if estilo(el, doc) in ({"No Spacing"} if os.environ.get("PETI_COMPACTO")
                                             else {"No Spacing", "Heading 1", "Heading 2"})}
    marcados = {id(el) for el in borrar}

    def solo_salto(el):
        if el.tag != qn("w:p") or id(el) in marcados:
            return False
        ppr = el.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            return False
        if any(es_decorativa(dr) for dr in el.iter(qn("w:drawing"))):
            return False
        return any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))

    for i, el in enumerate(hijos):
        if i == 0 or id(el) in usados or not solo_salto(el):
            continue
        siguiente = None
        for otro in hijos[i + 1:]:
            if id(otro) in marcados or (otro.tag == qn("w:p") and not texto_de(otro).strip()
                                         and not list(otro.iter(qn("w:drawing"))) and not solo_salto(otro)
                                         and (otro.find(qn("w:pPr")) is None
                                              or otro.find(qn("w:pPr")).find(qn("w:sectPr")) is None)):
                continue
            siguiente = otro
            break
        abre = siguiente is None or id(siguiente) in anclas_mayores or (
            siguiente.tag == qn("w:p") and siguiente.find(qn("w:pPr")) is not None
            and siguiente.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
        if not abre or solo_salto(siguiente):
            borrar.append(el); marcados.add(id(el))
    borrables = {id(el) for el in borrar}

    def punto_de_insercion(i, con_contenido):
        siguiente = next((p for p in posiciones if p > i), len(hijos))
        region = hijos[i + 1:siguiente]
        for el in region:
            if id(el) in borrables:
                return el
        ref = hijos[siguiente] if siguiente < len(hijos) else fin_cuerpo
        for el in region:
            ppr = el.find(qn("w:pPr")) if el.tag == qn("w:p") else None
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                ref = el
                break
        # Página separadora sin texto guía: el contenido empieza en la página siguiente
        decorada = any(es_decorativa(dr) for el in [hijos[i]] + region for dr in el.iter(qn("w:drawing")))
        if decorada and con_contenido:
            salto = parrafo_nuevo(); r = OxmlElement("w:r"); br = OxmlElement("w:br")
            br.set(qn("w:type"), "page"); r.append(br); salto.append(r)
            ref.addprevious(salto)
        return ref

    # Inserción del contenido
    ref, k, n_tabla, n_fig = None, -1, 0, 0
    for n, b in enumerate(bloques):
        if b[0] == "titulo":
            k += 1
            if k in asignados:
                i, el = asignados[k]
                if b[2] in RENOMBRAR:
                    # Solo se cambia el texto: se conservan el número, su formato y las figuras de la portadilla
                    ts = list(el.iter(qn("w:t")))
                    idx = next(k for k, t in enumerate(ts) if re.search(r"[A-Za-zÁÉÍÓÚáéíóú]", t.text or ""))
                    ts[idx].text = re.sub(r"^\d+\.\s*", "", b[2]) if idx else b[2]
                    for t in ts[idx + 1:]:
                        t.text = ""
                ref = punto_de_insercion(i, n + 1 < len(bloques) and bloques[n + 1][0] != "titulo")
            else:
                numerado = re.match(r"^\d+(\.\d+)+\s", b[2])
                # numerado -> Título 3 (como 6.1.x); sin número: capítulo nuevo (##) -> Título 2, subtítulo -> Título 4
                estilo_nuevo = "Heading 3" if numerado else ("Heading 2" if b[1] == 2 else "Heading 4")
                p = parrafo_nuevo(ids[estilo_nuevo])
                agregar_runs(p, b[2])
                ref.addprevious(p)
            continue
        if b[0] == "parrafo":
            p = parrafo_nuevo()
            agregar_runs(p, b[1]); ref.addprevious(p)
        elif b[0] == "vineta":
            p = parrafo_nuevo(ppr=vineta_ppr)
            agregar_runs(p, b[1]); ref.addprevious(p)
        elif b[0] == "imagen":
            _, desc, ruta = b
            p = parrafo_nuevo()
            ref.addprevious(p)
            par = Paragraph(p, doc._body)
            par.alignment = 1
            par.add_run().add_picture(os.path.join(RAIZ, ruta), width=Cm(16))
            if leyenda_fig is not None:
                n_fig += 1
                ref.addprevious(leyenda_con_numero(leyenda_fig, n_fig, desc))
        elif b[0] == "tabla":
            _, indice, desc, filas = b
            horizontal = max(len(f) for f in filas) >= HORIZONTAL
            if horizontal:
                # La sección que rige este punto termina en el siguiente salto de sección
                rige = next((s.find(qn("w:pPr")).find(qn("w:sectPr")) for s in ref.itersiblings()
                             if s.tag == qn("w:p") and s.find(qn("w:pPr")) is not None
                             and s.find(qn("w:pPr")).find(qn("w:sectPr")) is not None), fin_cuerpo)
                vertical = copy.deepcopy(rige)
                tipo = vertical.find(qn("w:type"))
                if tipo is not None:
                    vertical.remove(tipo)
                p = parrafo_nuevo(); ppr = OxmlElement("w:pPr"); ppr.append(vertical); p.append(ppr)
                ref.addprevious(p)
            if desc:
                n_tabla += 1
                ref.addprevious(leyenda_con_numero(leyenda, n_tabla, desc))
            ref.addprevious(clonar_tabla(modelos[indice], indice, filas, util_horizontal(rige) if horizontal else None))
            ref.addprevious(parrafo_nuevo())
            if horizontal:
                apaisada = copy.deepcopy(rige)
                tipo = apaisada.find(qn("w:type"))
                if tipo is not None:
                    apaisada.remove(tipo)
                pg = apaisada.find(qn("w:pgSz"))
                ancho, alto = pg.get(qn("w:w")), pg.get(qn("w:h"))
                pg.set(qn("w:w"), alto); pg.set(qn("w:h"), ancho); pg.set(qn("w:orient"), "landscape")
                ppr = OxmlElement("w:pPr"); ppr.append(apaisada)
                ultimo = ref.getprevious()
                ultimo.insert(0, ppr)
    for el in borrar:
        el.getparent().remove(el)
    if os.environ.get("PETI_COMPACTO"):
        # El glosario sigue al numeral 9 en la misma página: el salto de sección que lo antecede pasa a continuo
        cuerpo = list(doc.element.body)
        for i, el in enumerate(cuerpo):
            if el.tag == qn("w:p") and texto_de(el).strip() == "Glosario":
                for previo in reversed(cuerpo[max(0, i - 4):i]):
                    ppr = previo.find(qn("w:pPr")) if previo.tag == qn("w:p") else None
                    sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
                    if sect is not None:
                        tipo = sect.find(qn("w:type"))
                        if tipo is None:
                            tipo = OxmlElement("w:type")
                            sect.insert(0, tipo)
                        tipo.set(qn("w:val"), "continuous")
                        break
                break
        # Títulos que traen un salto de página después de su texto: el salto deja el título solo
        for p in doc.element.body.iter(qn("w:p")):
            if estilo(p, doc) != "Heading 2":
                continue
            runs = list(p.iter(qn("w:r")))
            ultimo_texto = max((i for i, r in enumerate(runs) if texto_de(r).strip()), default=None)
            if ultimo_texto is None:
                continue
            for r in runs[ultimo_texto + 1:]:
                for br in list(r.iter(qn("w:br"))):
                    if br.get(qn("w:type")) == "page":
                        br.getparent().remove(br)

    # Títulos sin morado: el estilo Título 4 de la plantilla (magenta) toma el naranja de los demás títulos
    from docx.shared import RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    t4 = doc.styles["Heading 4"]
    t4.font.color.rgb = RGBColor(0xE8, 0xA0, 0x43)
    t4.font.italic = False
    for par in doc.paragraphs:
        if par.style.name.startswith("Heading 4"):
            for run in par.runs:
                run.font.color.rgb = RGBColor(0xE8, 0xA0, 0x43)
                run.font.italic = False
        # Portadillas de capítulo ("7. Situación deseada u objetivo"): alineadas a la izquierda, no justificadas
        if par.style.name == "No Spacing" and 0 < len(par.text.strip()) < 70:
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.core_properties.title = "Plan Estratégico de Tecnologías de la Información (PETI) 2027-2030 — Valor + S.A.S."
    doc.core_properties.author = "Gerencia de Tecnología, Innovación y Sostenibilidad"
    doc.save(salida)
    print("Generado:", salida)
    print("Títulos de la plantilla usados:", len(asignados), "| títulos nuevos:", len(nuevos), "| tablas:", n_tabla)
    for t in nuevos:
        print("  nuevo:", t)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    convertir(sys.argv[1], sys.argv[2])
